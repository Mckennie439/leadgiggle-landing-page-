import os

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
import datetime

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Referral_Program_CAC_Tracker.xlsx")

# ---------- palette ----------
DARK_RED = "C8272A"
WHITE = "FFFFFF"
YELLOW = "FFFF99"
LIGHT_GRAY = "F2F2F2"
BORDER_GRAY = "BFBFBF"
BLUE_INPUT = "0000FF"
GREEN = "C6EFCE"
GREEN_TEXT = "006100"
YEL_FILL = "FFEB9C"
YEL_TEXT = "9C6500"
RED_FILL = "FFC7CE"
RED_TEXT = "9C0006"

FONT_NAME = "Arial"

thin = Side(style="thin", color=BORDER_GRAY)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

header_font = Font(name=FONT_NAME, bold=True, color=WHITE, size=10)
header_fill = PatternFill("solid", fgColor=DARK_RED)
title_font = Font(name=FONT_NAME, bold=True, color=WHITE, size=14)
subheader_font = Font(name=FONT_NAME, bold=True, color=WHITE, size=10)
label_font = Font(name=FONT_NAME, size=10)
label_bold = Font(name=FONT_NAME, size=10, bold=True)
value_font = Font(name=FONT_NAME, size=10)
input_font = Font(name=FONT_NAME, size=10, color=BLUE_INPUT)
formula_font = Font(name=FONT_NAME, size=10, color="000000")
note_font = Font(name=FONT_NAME, size=9, italic=True, color="808080")
yellow_fill = PatternFill("solid", fgColor=YELLOW)

CURRENCY_FMT = '"KES" #,##0'
PCT_FMT = '0.0%'
DATE_FMT = 'DD-MMM-YYYY'

wb = Workbook()

# =========================================================================
# SHEET 1: Install Log
# =========================================================================
ws = wb.active
ws.title = "Install Log"

HEADERS = [
    "Install #", "Referral Code", "Referrer Name", "Installation Date",
    "Installation Value (KES)", "Tier", "Commission %", "Commission Paid (KES)",
    "Discount %", "Discount Given (KES)", "Cost of Materials/Labor/Logistics (KES)",
    "Gross Margin (KES)", "Gross Margin %", "Overhead Allocation (KES)",
    "Total Acquisition Cost (KES)", "CAC % of Deal Value", "CAC-to-Margin Ratio", "Flag",
]
N_COLS = len(HEADERS)  # 18 -> A..R

COL_WIDTHS = [10, 16, 20, 16, 20, 14, 13, 18, 11, 17, 26, 16, 14, 20, 20, 16, 16, 12]
for i, w in enumerate(COL_WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# --- Title row ---
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N_COLS)
c = ws.cell(row=1, column=1, value="Referral Program CAC Tracker — Install Log")
c.font = title_font
c.fill = header_fill
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[1].height = 24

# --- Legend row ---
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=N_COLS)
legend = ws.cell(
    row=2, column=1,
    value=("Legend:  Blue text = enter manually.  Yellow fill = editable rate assumptions "
           "(update after the validation phase).  Black text = auto-calculated — do not edit."),
)
legend.font = note_font
legend.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[2].height = 16

# --- Rate settings block (rows 3-5) ---
# Labels are merged across A:C so the text has room; column A alone stays
# narrow (10 chars) because it doubles as the "Install #" column below.
ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=3)
ws.cell(row=3, column=1, value="Commission Rate (all tiers):").font = label_bold
rc = ws.cell(row=3, column=4, value=0.05)
rc.font = value_font
rc.fill = yellow_fill
rc.number_format = PCT_FMT
rc.border = box

ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=3)
ws.cell(row=4, column=1, value="Discount Rate – Above 250k:").font = label_bold
rd1 = ws.cell(row=4, column=4, value=0.04)
rd1.font = value_font
rd1.fill = yellow_fill
rd1.number_format = PCT_FMT
rd1.border = box

ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=3)
ws.cell(row=5, column=1, value="Discount Rate – 150k–250k:").font = label_bold
rd2 = ws.cell(row=5, column=4, value=0.03)
rd2.font = value_font
rd2.fill = yellow_fill
rd2.number_format = PCT_FMT
rd2.border = box

COMM_CELL = "$D$3"
DISC_ABOVE_CELL = "$D$4"
DISC_MID_CELL = "$D$5"

# --- Table header (row 7) ---
HEADER_ROW = 7
for i, h in enumerate(HEADERS, start=1):
    cell = ws.cell(row=HEADER_ROW, column=i, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = box
ws.row_dimensions[HEADER_ROW].height = 32

FIRST_DATA_ROW = 8
N_INSTALLS = 30
LAST_DATA_ROW = FIRST_DATA_ROW + N_INSTALLS - 1  # 37

MANUAL_COLS = {2, 3, 4, 5, 11, 14}  # Referral Code, Referrer Name, Date, Install Value, Cost, Overhead

# Example row (Install #1) realistic sample values
EXAMPLE = {
    2: "MF-0001",
    3: "Sample Referrer",
    4: datetime.date(2026, 1, 15),
    5: 275000,
    11: 190000,
    14: 3000,
}

for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
    # A: Install #
    a = ws.cell(row=r, column=1, value=f"=ROW()-{FIRST_DATA_ROW - 1}")
    a.font = formula_font
    a.alignment = Alignment(horizontal="center")

    # B-E, K, N: manual entry columns
    for col in (2, 3, 4, 5, 11, 14):
        cell = ws.cell(row=r, column=col)
        if r == FIRST_DATA_ROW and col in EXAMPLE:
            cell.value = EXAMPLE[col]
        cell.font = input_font

    ws.cell(row=r, column=4).number_format = DATE_FMT
    ws.cell(row=r, column=5).number_format = CURRENCY_FMT
    ws.cell(row=r, column=11).number_format = CURRENCY_FMT
    ws.cell(row=r, column=14).number_format = CURRENCY_FMT

    # F: Tier
    f = ws.cell(
        row=r, column=6,
        value=f'=IF(E{r}="","",IF(E{r}>=250000,"Above 250k",IF(E{r}>=150000,"150k-250k","Below 150k")))',
    )
    f.font = formula_font

    # G: Commission %
    g = ws.cell(row=r, column=7, value=f'=IF(E{r}="","",{COMM_CELL})')
    g.font = formula_font
    g.number_format = PCT_FMT

    # H: Commission Paid (KES)
    h = ws.cell(row=r, column=8, value=f'=IF(E{r}="","",E{r}*G{r})')
    h.font = formula_font
    h.number_format = CURRENCY_FMT

    # I: Discount %
    i_ = ws.cell(
        row=r, column=9,
        value=f'=IF(F{r}="","",IF(F{r}="Above 250k",{DISC_ABOVE_CELL},IF(F{r}="150k-250k",{DISC_MID_CELL},0)))',
    )
    i_.font = formula_font
    i_.number_format = PCT_FMT

    # J: Discount Given (KES)
    j = ws.cell(row=r, column=10, value=f'=IF(E{r}="","",E{r}*I{r})')
    j.font = formula_font
    j.number_format = CURRENCY_FMT

    # L: Gross Margin (KES)
    l = ws.cell(row=r, column=12, value=f'=IF(OR(E{r}="",K{r}=""),"",E{r}-K{r})')
    l.font = formula_font
    l.number_format = CURRENCY_FMT

    # M: Gross Margin %
    m = ws.cell(row=r, column=13, value=f'=IF(OR(L{r}="",E{r}="",E{r}=0),"",L{r}/E{r})')
    m.font = formula_font
    m.number_format = PCT_FMT

    # O: Total Acquisition Cost (KES)
    o = ws.cell(row=r, column=15, value=f'=IF(OR(E{r}="",N{r}=""),"",H{r}+J{r}+N{r})')
    o.font = formula_font
    o.number_format = CURRENCY_FMT

    # P: CAC % of Deal Value
    p = ws.cell(row=r, column=16, value=f'=IF(OR(O{r}="",E{r}="",E{r}=0),"",O{r}/E{r})')
    p.font = formula_font
    p.number_format = PCT_FMT

    # Q: CAC-to-Margin Ratio
    q = ws.cell(row=r, column=17, value=f'=IF(OR(O{r}="",L{r}="",L{r}=0),"",O{r}/L{r})')
    q.font = formula_font
    q.number_format = '0.00'

    # R: Flag
    flag = ws.cell(
        row=r, column=18,
        value=f'=IF(Q{r}="","",IF(Q{r}>0.4,"High Risk",IF(Q{r}>=0.3,"Caution","Healthy")))',
    )
    flag.font = Font(name=FONT_NAME, size=10, bold=True)
    flag.alignment = Alignment(horizontal="center")

    for col in range(1, N_COLS + 1):
        ws.cell(row=r, column=col).border = box

ws.freeze_panes = "A8"

# --- Conditional formatting on Flag column ---
flag_range = f"R{FIRST_DATA_ROW}:R{LAST_DATA_ROW}"
q_anchor = f"$Q{FIRST_DATA_ROW}"
ws.conditional_formatting.add(
    flag_range,
    FormulaRule(formula=[f'AND({q_anchor}<>"",{q_anchor}>0.4)'],
                fill=PatternFill("solid", fgColor=RED_FILL), font=Font(color=RED_TEXT, bold=True)),
)
ws.conditional_formatting.add(
    flag_range,
    FormulaRule(formula=[f'AND({q_anchor}<>"",{q_anchor}>=0.3,{q_anchor}<=0.4)'],
                fill=PatternFill("solid", fgColor=YEL_FILL), font=Font(color=YEL_TEXT, bold=True)),
)
ws.conditional_formatting.add(
    flag_range,
    FormulaRule(formula=[f'AND({q_anchor}<>"",{q_anchor}<0.3)'],
                fill=PatternFill("solid", fgColor=GREEN), font=Font(color=GREEN_TEXT, bold=True)),
)

# ---- Footnote below table documenting the tier assumption ----
note_row = LAST_DATA_ROW + 2
ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=N_COLS)
note = ws.cell(
    row=note_row, column=1,
    value=("Note: Tier and rate thresholds reflect the referral program structure as of "
           "2026-08-06 (source: program terms provided by G&P Solar Systems). "
           "Installs below KES 150,000 are logged with Tier = \"Below 150k\" and a 0% discount rate "
           "pending a decision on whether they qualify for the program."),
)
note.font = note_font
note.alignment = Alignment(horizontal="left", wrap_text=True)

# =========================================================================
# SHEET 2: Summary Dashboard
# =========================================================================
ws2 = wb.create_sheet("Summary Dashboard")
for i, w in enumerate([34, 20, 14, 14], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w

ws2.merge_cells("A1:D1")
t = ws2.cell(row=1, column=1, value="Referral Program CAC Tracker — Summary Dashboard")
t.font = title_font
t.fill = header_fill
t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws2.row_dimensions[1].height = 24

LOG = "'Install Log'"
E_RNG = f"{LOG}!E{FIRST_DATA_ROW}:E{LAST_DATA_ROW}"
H_RNG = f"{LOG}!H{FIRST_DATA_ROW}:H{LAST_DATA_ROW}"
J_RNG = f"{LOG}!J{FIRST_DATA_ROW}:J{LAST_DATA_ROW}"
N_RNG = f"{LOG}!N{FIRST_DATA_ROW}:N{LAST_DATA_ROW}"
O_RNG = f"{LOG}!O{FIRST_DATA_ROW}:O{LAST_DATA_ROW}"
M_RNG = f"{LOG}!M{FIRST_DATA_ROW}:M{LAST_DATA_ROW}"
Q_RNG = f"{LOG}!Q{FIRST_DATA_ROW}:Q{LAST_DATA_ROW}"
HALF1_ROWS = (FIRST_DATA_ROW, FIRST_DATA_ROW + 14)          # rows 8-22 -> installs 1-15
HALF2_ROWS = (FIRST_DATA_ROW + 15, LAST_DATA_ROW)            # rows 23-37 -> installs 16-30
Q_HALF1 = f"{LOG}!Q{HALF1_ROWS[0]}:Q{HALF1_ROWS[1]}"
Q_HALF2 = f"{LOG}!Q{HALF2_ROWS[0]}:Q{HALF2_ROWS[1]}"


def label_value(row, label, formula, fmt=None, bold_label=False):
    lc = ws2.cell(row=row, column=1, value=label)
    lc.font = label_bold if bold_label else label_font
    vc = ws2.cell(row=row, column=2, value=formula)
    vc.font = value_font
    if fmt:
        vc.number_format = fmt
    lc.border = box
    vc.border = box
    return vc


def section_header(row, text, span=2):
    ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws2.cell(row=row, column=1, value=text)
    c.font = subheader_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return c


# Totals
r = 3
v = label_value(r, "Installs Logged", f"=COUNT({E_RNG})")
v.number_format = "0"
ws2.cell(row=r, column=3, value="of 30").font = note_font
r += 1
label_value(r, "Total Installation Value (KES)", f"=SUM({E_RNG})", CURRENCY_FMT); r += 1
label_value(r, "Total Commissions Paid (KES)", f"=SUM({H_RNG})", CURRENCY_FMT); r += 1
label_value(r, "Total Discounts Given (KES)", f"=SUM({J_RNG})", CURRENCY_FMT); r += 1
label_value(r, "Total Overhead Allocated (KES)", f"=SUM({N_RNG})", CURRENCY_FMT); r += 1

r += 1
section_header(r, "Blended CAC Metrics"); r += 1
avg_deal_row = r
label_value(r, "Average Deal Value (KES)",
            f'=IF(COUNT({E_RNG})=0,"",SUM({E_RNG})/COUNT({E_RNG}))', CURRENCY_FMT); r += 1
blended_cac_row = r
label_value(r, "Blended CAC (KES)",
            f'=IF(COUNT({E_RNG})=0,"",SUM({O_RNG})/COUNT({E_RNG}))', CURRENCY_FMT); r += 1
label_value(r, "Blended CAC % of Avg Deal Value",
            f'=IF(OR(B{avg_deal_row}="",B{avg_deal_row}=0),"",B{blended_cac_row}/B{avg_deal_row})', PCT_FMT); r += 1
avg_margin_row = r
label_value(r, "Average Gross Margin %", f"=IFERROR(AVERAGE({M_RNG}),\"\")", PCT_FMT); r += 1
avg_ratio_row = r
label_value(r, "Average CAC-to-Margin Ratio", f"=IFERROR(AVERAGE({Q_RNG}),\"\")", '0.00'); r += 1

r += 1
section_header(r, "Trend Check — Installs 1–15 vs 16–30"); r += 1
half1_row = r
label_value(r, "Avg CAC-to-Margin Ratio, Installs 1–15", f"=IFERROR(AVERAGE({Q_HALF1}),\"\")", '0.00'); r += 1
half2_row = r
label_value(r, "Avg CAC-to-Margin Ratio, Installs 16–30", f"=IFERROR(AVERAGE({Q_HALF2}),\"\")", '0.00'); r += 1
trend_row = r
label_value(
    r, "Trend Indicator",
    (f'=IF(OR(B{half1_row}="",B{half2_row}=""),"Insufficient Data",'
     f'IF(B{half2_row}>B{half1_row}*1.05,"Climbing",'
     f'IF(B{half2_row}<B{half1_row}*0.95,"Improving","Stable")))'),
    bold_label=True,
)
ws2.cell(row=r, column=2).font = Font(name=FONT_NAME, size=10, bold=True)
r += 1
ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
tn = ws2.cell(
    row=r, column=1,
    value=('Assumption: "Climbing" = 2nd-half average ratio more than 5% above the 1st-half '
           'average; "Improving" = more than 5% below. Adjust the 1.05 / 0.95 multipliers in the '
           'formula if a different sensitivity is wanted.'),
)
tn.font = note_font
tn.alignment = Alignment(wrap_text=True)
r += 2

section_header(r, "Program Status", span=4); r += 1
status_row = r
ws2.merge_cells(start_row=status_row, start_column=1, end_row=status_row + 3, end_column=4)
status_cell = ws2.cell(
    row=status_row, column=1,
    value=(f'=IF(B{avg_ratio_row}="","INSUFFICIENT DATA",'
           f'IF(B{avg_ratio_row}>0.4,"RENEGOTIATE TERMS",'
           f'IF(B{avg_ratio_row}>=0.35,"REVIEW NEEDED",'
           f'IF(B{trend_row}="Climbing","REVIEW NEEDED","SAFE TO SCALE"))))'),
)
status_cell.font = Font(name=FONT_NAME, size=26, bold=True, color=WHITE)
status_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
status_cell.fill = PatternFill("solid", fgColor="808080")
for rr in range(status_row, status_row + 4):
    for cc in range(1, 5):
        ws2.cell(row=rr, column=cc).border = box
ws2.row_dimensions[status_row].height = 20
for rr in range(status_row + 1, status_row + 4):
    ws2.row_dimensions[rr].height = 20

status_anchor = f"$A${status_row}"
ws2.conditional_formatting.add(
    f"A{status_row}:D{status_row + 3}",
    FormulaRule(formula=[f'{status_anchor}="RENEGOTIATE TERMS"'],
                fill=PatternFill("solid", fgColor="C0392B")),
)
ws2.conditional_formatting.add(
    f"A{status_row}:D{status_row + 3}",
    FormulaRule(formula=[f'{status_anchor}="REVIEW NEEDED"'],
                fill=PatternFill("solid", fgColor="D4AC0D")),
)
ws2.conditional_formatting.add(
    f"A{status_row}:D{status_row + 3}",
    FormulaRule(formula=[f'{status_anchor}="SAFE TO SCALE"'],
                fill=PatternFill("solid", fgColor="1E8449")),
)

r = status_row + 5
ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
sn = ws2.cell(
    row=r, column=1,
    value=('Thresholds: RENEGOTIATE TERMS if Average CAC-to-Margin Ratio > 0.40 · '
           'REVIEW NEEDED if 0.35–0.40, or under 0.35 while the trend is "Climbing" · '
           'SAFE TO SCALE if under 0.35 and the trend is not "Climbing".'),
)
sn.font = note_font
sn.alignment = Alignment(wrap_text=True)

ws2.freeze_panes = "A2"

wb.save(OUT)
print("saved", OUT)
